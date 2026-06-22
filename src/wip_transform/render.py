"""Phase 4 — 輸出 Excel (render).

寫出報表(16 欄、固定順序),反紅列套用紅色格式。讀寫與轉換分離:
這支只負責把 transform() 的結果寫成 .xlsx。
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .transform import (
    FLAG_LOTTYPE,
    FLAG_STAYTIME,
    LOTTYPE_COLUMN,
    OUTPUT_COLUMNS,
    STAYTIME_COLUMN,
)

# 字紅 (red font) for any flagged row; 底紅 (red fill) on the specific triggering cell.
_RED_FONT = Font(color="FFCC0000")
_RED_FILL = PatternFill(fill_type="solid", fgColor="FFF4CCCC")


class RenderError(RuntimeError):
    """Raised when the Phase 4 VERIFY (red-row count) fails."""


def write_report(df_out: pd.DataFrame, path: str) -> int:
    """Write the 16-column report to `path`, red-flagging hit rows.

    Returns the number of red rows (rows hit by rule A or B). VERIFY: that
    count must equal the rows transform() flagged.
    """
    red_mask = (df_out[FLAG_STAYTIME] | df_out[FLAG_LOTTYPE]).to_numpy()
    expected_red = int(red_mask.sum())

    wb = Workbook()
    ws = wb.active
    ws.title = "WIP"

    # header (fixed 16-column order)
    ws.append(OUTPUT_COLUMNS)

    staytime_col_idx = OUTPUT_COLUMNS.index(STAYTIME_COLUMN) + 1
    lottype_col_idx = OUTPUT_COLUMNS.index(LOTTYPE_COLUMN) + 1

    rendered_red = 0
    for _, row in df_out.iterrows():
        ws.append([row[c] for c in OUTPUT_COLUMNS])
        excel_row = ws.max_row

        is_red = bool(row[FLAG_STAYTIME] or row[FLAG_LOTTYPE])
        if is_red:
            rendered_red += 1
            for col_idx in range(1, len(OUTPUT_COLUMNS) + 1):
                ws.cell(row=excel_row, column=col_idx).font = _RED_FONT
        if bool(row[FLAG_STAYTIME]):
            ws.cell(row=excel_row, column=staytime_col_idx).fill = _RED_FILL
        if bool(row[FLAG_LOTTYPE]):
            ws.cell(row=excel_row, column=lottype_col_idx).fill = _RED_FILL

    # VERIFY (§ Phase 4): rendered red rows == flagged rows.
    if rendered_red != expected_red:
        raise RenderError(
            f"反紅列數不符:預期 {expected_red},實際渲染 {rendered_red}。停止。"
        )

    wb.save(path)
    return rendered_red
